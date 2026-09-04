"""Uploaded file bytes -> plain text.

Scope is deliberately narrow, per the architecture spec: typed source material
only. OCR of scanned pages or handwritten notes is explicitly out of scope —
an organisation with only scanned pages pastes the transcribed text instead.
That single cut is what keeps ingestion buildable.
"""

from __future__ import annotations

import io
from pathlib import Path

from fastapi import HTTPException, status

SUPPORTED_SUFFIXES = {".txt", ".md", ".csv", ".vtt", ".srt", ".docx", ".pdf", ".pptx"}


class ExtractionError(Exception):
    """Raised when a file cannot be turned into usable text."""


def _decode_plain(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("The file is not readable as text in any common encoding.")


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ExtractionError("python-docx is not installed.") from exc

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    # Session notes are frequently tabular (attendance, per-district figures),
    # and the paragraph walk above skips tables entirely.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ExtractionError("pypdf is not installed.") from exc

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(p for p in pages if p.strip())

    if not text.strip():
        # A PDF of scanned images extracts to nothing. Say so precisely rather
        # than passing an empty string down the pipeline and reporting that
        # zero identifiers were found in it.
        raise ExtractionError(
            "No text layer found. This looks like a scanned PDF, and OCR is not "
            "supported — paste the transcribed text instead."
        )
    return text


def _extract_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ExtractionError("python-pptx is not installed.") from exc

    presentation = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        slide_parts = [
            shape.text.strip()
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        ]
        if slide_parts:
            parts.append(f"[Slide {index}]\n" + "\n".join(slide_parts))
    return "\n\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """Dispatch on file extension and return normalised plain text."""
    suffix = Path(filename or "").suffix.lower()

    if suffix not in SUPPORTED_SUFFIXES:
        raise ExtractionError(
            f"Unsupported file type '{suffix or 'unknown'}'. Supported: "
            + ", ".join(sorted(SUPPORTED_SUFFIXES))
            + ". Scanned images and handwriting are not supported."
        )

    if suffix == ".docx":
        text = _extract_docx(data)
    elif suffix == ".pdf":
        text = _extract_pdf(data)
    elif suffix == ".pptx":
        text = _extract_pptx(data)
    else:
        text = _decode_plain(data)

    normalised = "\n".join(line.rstrip() for line in text.splitlines())
    normalised = "\n".join(filter(None, (ln.strip() and ln for ln in normalised.splitlines())))

    if not normalised.strip():
        raise ExtractionError("The file contained no extractable text.")

    return normalised.strip()


def extraction_http_error(exc: ExtractionError) -> HTTPException:
    return HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(exc))
